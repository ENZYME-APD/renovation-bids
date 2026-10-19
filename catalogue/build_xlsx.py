"""Build the Enzyme catalogue workbook from es_catalogue.py.

Usage: python catalogue/build_xlsx.py  ->  catalogue/out/Enzyme_catalogo_partidas_ES.xlsx
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import es_catalogue as C

OUT = Path(__file__).parent / "out" / "Enzyme_catalogo_partidas_ES.xlsx"

F = "Arial"
BASE = Font(name=F, size=10)
BOLD = Font(name=F, size=10, bold=True)
HEAD = Font(name=F, size=10, bold=True, color="FFFFFF")
TITLE = Font(name=F, size=14, bold=True)
INPUT = Font(name=F, size=10, color="0000FF")
LINK = Font(name=F, size=10, color="008000")
HEAD_FILL = PatternFill("solid", fgColor="1C2321")
CHAP_FILL = PatternFill("solid", fgColor="E8E3D8")
YELLOW = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="D9D3C7")
BOX = Border(bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
EUR = '#,##0.00 €;-#,##0.00 €;"-"'
EUR0 = '#,##0 €;-#,##0 €;"-"'
QTY = '#,##0.00;-#,##0.00;"-"'


def header(ws, row, labels, widths=None):
    for i, label in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=label)
        c.font, c.fill = HEAD, HEAD_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def style_all(ws):
    for row in ws.iter_rows():
        for c in row:
            if c.font == Font():  # untouched default
                c.font = BASE


wb = Workbook()

# ── Léeme ─────────────────────────────────────────────────────────────────
ws = wb.active
ws.title = "Léeme"
ws.column_dimensions["A"].width = 110
lines = [
    ("Catálogo de partidas Enzyme — Reforma de vivienda (España)", TITLE),
    ("Versión 1 · base de precios: presupuesto Ortega y Gasset 47 (P26018, REV 2) · precios en € sin IVA", BASE),
    ("", BASE),
    ("Cómo está organizado", BOLD),
    ("• Capítulos por gremio, cada uno referido a un capítulo del Generador de Precios de CYPE (Rehabilitación).", BASE),
    ("• Un código = una partida con precio unitario. Los códigos no cambian de significado; una especificación nueva lleva código nuevo.", BASE),
    ("• Todas las partidas son 'todo incluido' (suministro + colocación) en la calidad estándar Enzyme.", BASE),
    ("• Fase de oferta: solo estándar, sin opciones ni peticiones del cliente. Las alternativas (tarima, PVC, gran altura…) son partidas del catálogo para después del contrato.", BASE),
    ("", BASE),
    ("Hojas", BOLD),
    ("• Capítulos: estructura y correspondencia con CYPE.", BASE),
    ("• Partidas: catálogo con precio Enzyme y de dónde sale cada precio.", BASE),
    ("• Plantillas: estancias tipo. Cantidad = factor × tamaño + fijo.", BASE),
    ("• Calculadora: estancias y m² → presupuesto por capítulos. Ejemplo con superficies aproximadas; ajústelas.", BASE),
    ("• Las plantillas cubren lo que depende de la superficie y del tipo de estancia. Se añaden a mano las partidas de proyecto: tabiquería y su demolición, armarios y carpintería a medida, calefacción y aire acondicionado, molduras.", BASE),
    ("", BASE),
    ("Colores", BOLD),
    ("• Azul: dato introducido (precios, factores, estancias, m²). Negro: fórmula. Verde: dato traído de otra hoja.", BASE),
    ("• Fondo amarillo: pendiente de completar (precio o referencia CYPE).", BASE),
    ("", BASE),
    ("Qué falta", BOLD),
    ("• Partidas sin precio: nuevas respecto a Ortega y Gasset 47. Precio por CYPE o por solicitud a proveedor.", BASE),
    ("• Ref. CYPE: código de la partida equivalente del Generador de Precios, a rellenar al tarar precios.", BASE),
    ("• Factores de plantillas sin nota 'OG47': valores propuestos, a calibrar con 2–3 obras terminadas.", BASE),
    ("", BASE),
    ("Partidas de Ortega y Gasset 47 que no pasan al catálogo", BOLD),
]
for text, font in lines:
    ws.append([text])
    ws.cell(row=ws.max_row, column=1).font = font
for item, why in C.OG47_EXCLUDED:
    ws.append([f"• {item}: {why}"])
    ws.cell(row=ws.max_row, column=1).font = BASE
for r in ws.iter_rows():
    r[0].alignment = Alignment(wrap_text=True, vertical="top")

# ── Capítulos ─────────────────────────────────────────────────────────────
ws = wb.create_sheet("Capítulos")
header(ws, 1, ["Código", "Capítulo", "Capítulo CYPE", "Orden", "Nº partidas"], [10, 48, 14, 8, 12])
for code, name, cype, seq in C.CHAPTERS:
    r = ws.max_row + 1
    ws.append([code, name, cype, seq, f'=COUNTIF(Partidas!$B:$B,A{r})'])
style_all(ws)

# ── Partidas ──────────────────────────────────────────────────────────────
ws = wb.create_sheet("Partidas")
cols = ["Código", "Capítulo", "Ud", "Descripción", "Especificación estándar Enzyme",
        "Precio Enzyme (€/ud)", "Base del precio", "Ref. CYPE", "Estado"]
header(ws, 1, cols, [11, 9, 6, 42, 70, 14, 46, 12, 12])
chapter_name = {c[0]: c[1] for c in C.CHAPTERS}
first = 2
for p in C.PARTIDAS:
    r = ws.max_row + 1
    ws.append([p["code"], p["code"][:3], p["unit"], p["short"], p["spec"], p["price"], p["basis"], None,
               f'=IF(F{r}="","Sin precio","Con precio")'])
    ws.cell(row=r, column=6).font = INPUT
    ws.cell(row=r, column=6).number_format = EUR
    for col in (4, 5, 7):
        ws.cell(row=r, column=col).alignment = WRAP
    if p["price"] is None:
        ws.cell(row=r, column=6).fill = YELLOW
    ws.cell(row=r, column=8).fill = YELLOW
last_p = ws.max_row
ws.auto_filter.ref = f"A1:I{last_p}"
ws.cell(row=1, column=8).comment = Comment("Código de la partida equivalente en el Generador de Precios de CYPE (Rehabilitación). A rellenar al tarar precios.", "Enzyme")
ws.cell(row=1, column=6).comment = Comment("€ sin IVA, todo incluido (suministro + colocación), calidad estándar Enzyme. Fuente en 'Base del precio'.", "Enzyme")
style_all(ws)
PR = f"Partidas!$A$2:$A${last_p}"

# ── Plantillas ────────────────────────────────────────────────────────────
ws = wb.create_sheet("Plantillas")
header(ws, 1, ["Plantilla", "Estancia", "Tamaño en", "Partida", "Descripción", "Ud", "Factor", "Fijo", "Nota"],
       [10, 30, 22, 11, 46, 6, 9, 9, 44])
for t in C.TEMPLATES:
    for ln in t["lines"]:
        r = ws.max_row + 1
        ws.append([t["code"], t["name"], t["size_label"], ln["partida"],
                   f'=INDEX(Partidas!$D$2:$D${last_p},MATCH(D{r},{PR},0))',
                   f'=INDEX(Partidas!$C$2:$C${last_p},MATCH(D{r},{PR},0))',
                   ln["factor"], ln["fixed"], ln["note"]])
        for col in (5, 6):
            ws.cell(row=r, column=col).font = LINK
        for col in (7, 8):
            ws.cell(row=r, column=col).font = INPUT
            ws.cell(row=r, column=col).number_format = '0.000;-0.000;"-"'
last_t = ws.max_row
ws.auto_filter.ref = f"A1:I{last_t}"
style_all(ws)

# ── Calculadora ───────────────────────────────────────────────────────────
ws = wb.create_sheet("Calculadora")
ROOMS = 14
example = [  # name, template, size  — approximate split of Ortega y Gasset 47 (91 m²)
    ("Vivienda", "GEN", 91), ("Salón-comedor", "SAL", 24), ("Dormitorio principal", "DOR", 15),
    ("Dormitorio 2", "DOR", 10), ("Dormitorio 3", "DOR", 10), ("Pasillo", "PAS", 7), ("Hall", "PAS", 4),
    ("Cocina", "COC", 9), ("Lavandería", "LAV", 3), ("Baño principal", "BAN", 5), ("Baño 2", "BAN", 4),
    ("Ventanas", "VEN", 9),
]
ws["A1"] = "Calculadora de oferta — calidad estándar Enzyme"
ws["A1"].font = TITLE
ws["A2"] = "Rellene en azul: nombre de estancia, plantilla y tamaño (m², o nº de ventanas para VEN). Ejemplo con superficies aproximadas de Ortega y Gasset 47; ajústelas."
ws["A2"].font = BASE
ws["A4"], ws["A5"], ws["A6"] = "Estancia", "Plantilla", "Tamaño"
for c in ("A4", "A5", "A6"):
    ws[c].font = BOLD
FIRST_ROOM_COL = 5  # E
dv = DataValidation(type="list", formula1=f'"{",".join(t["code"] for t in C.TEMPLATES)}"', allow_blank=True)
ws.add_data_validation(dv)
for j in range(ROOMS):
    col = FIRST_ROOM_COL + j
    L = get_column_letter(col)
    ws.column_dimensions[L].width = 11
    name, tpl, size = example[j] if j < len(example) else (None, None, None)
    for row, val in ((4, name), (5, tpl), (6, size)):
        c = ws.cell(row=row, column=col, value=val)
        c.font = INPUT
        c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="center")
    dv.add(f"{L}5")
ws.row_dimensions[4].height = 30

HROW = 8
tot_col = FIRST_ROOM_COL + ROOMS          # total quantity
price_col, amount_col = tot_col + 1, tot_col + 2
labels = ["Código", "Cap.", "Descripción", "Ud"] + [""] * ROOMS + ["Cantidad", "Precio (€/ud)", "Importe (€)"]
for i, label in enumerate(labels, 1):
    c = ws.cell(row=HROW, column=i, value=label)
    c.font, c.fill = HEAD, HEAD_FILL
for j in range(ROOMS):
    L = get_column_letter(FIRST_ROOM_COL + j)
    ws.cell(row=HROW, column=FIRST_ROOM_COL + j, value=f"={L}4&\"\"")
ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 10, 6
ws.column_dimensions["C"].width, ws.column_dimensions["D"].width = 44, 5
for col in (tot_col, price_col, amount_col):
    ws.column_dimensions[get_column_letter(col)].width = 13

TA, TD, TG, TH = (f"Plantillas!${c}$2:${c}${last_t}" for c in "ADGH")
r = HROW + 1
first_line = r
for code, cname, _cype, _seq in C.CHAPTERS:
    ws.cell(row=r, column=1, value=code).font = BOLD
    ws.cell(row=r, column=3, value=cname).font = BOLD
    for col in range(1, amount_col + 1):
        ws.cell(row=r, column=col).fill = CHAP_FILL
    chap_row = r
    r += 1
    start = r
    for p in [p for p in C.PARTIDAS if p["code"].startswith(code)]:
        ws.cell(row=r, column=1, value=p["code"])
        ws.cell(row=r, column=2, value=code)
        ws.cell(row=r, column=3, value=f'=INDEX(Partidas!$D$2:$D${last_p},MATCH(A{r},{PR},0))').font = LINK
        ws.cell(row=r, column=4, value=f'=INDEX(Partidas!$C$2:$C${last_p},MATCH(A{r},{PR},0))').font = LINK
        for j in range(ROOMS):
            L = get_column_letter(FIRST_ROOM_COL + j)
            f = (f"=IF({L}$5=\"\",0,SUMPRODUCT(({TA}={L}$5)*({TD}=$A{r})*({TG}*N({L}$6)+{TH})))")
            c = ws.cell(row=r, column=FIRST_ROOM_COL + j, value=f)
            c.number_format = QTY
        a, b = get_column_letter(FIRST_ROOM_COL), get_column_letter(FIRST_ROOM_COL + ROOMS - 1)
        ws.cell(row=r, column=tot_col, value=f"=SUM({a}{r}:{b}{r})").number_format = QTY
        pc = ws.cell(row=r, column=price_col,
                     value=f'=IF(INDEX(Partidas!$F$2:$F${last_p},MATCH(A{r},{PR},0))="","",INDEX(Partidas!$F$2:$F${last_p},MATCH(A{r},{PR},0)))')
        pc.font, pc.number_format = LINK, EUR
        P_, Q_ = get_column_letter(price_col), get_column_letter(tot_col)
        ws.cell(row=r, column=amount_col, value=f'=IF({P_}{r}="",0,{Q_}{r}*{P_}{r})').number_format = EUR
        r += 1
    A_ = get_column_letter(amount_col)
    c = ws.cell(row=chap_row, column=amount_col, value=f"=SUM({A_}{start}:{A_}{r - 1})")
    c.font, c.number_format = BOLD, EUR0
last_line = r - 1

A_, Q_, P_ = get_column_letter(amount_col), get_column_letter(tot_col), get_column_letter(price_col)
r += 1
summary = [
    ("Ejecución material (partidas con precio)", f'=SUMIF($B${first_line}:$B${last_line},"<>",{A_}{first_line}:{A_}{last_line})', EUR0, BOLD),
    ("Imprevistos (%)", 0.07, "0.0%", INPUT),
    ("Imprevistos", None, EUR0, BASE),
    ("Total obra sin IVA", None, EUR0, BOLD),
    ("IVA (%)", 0.10, "0.0%", INPUT),
    ("Total obra con IVA", None, EUR0, BOLD),
    ("Partidas con cantidad y sin precio", f'=SUMPRODUCT(({Q_}{first_line}:{Q_}{last_line}>0)*({P_}{first_line}:{P_}{last_line}=""))', "0", BOLD),
]
rows = {}
for label, val, fmt, font in summary:
    ws.cell(row=r, column=3, value=label).font = BOLD if font is BOLD else BASE
    c = ws.cell(row=r, column=amount_col, value=val)
    c.number_format, c.font = fmt, font
    rows[label] = r
    r += 1
ws.cell(row=rows["Imprevistos"], column=amount_col,
        value=f"={A_}{rows['Ejecución material (partidas con precio)']}*{A_}{rows['Imprevistos (%)']}")
ws.cell(row=rows["Total obra sin IVA"], column=amount_col,
        value=f"={A_}{rows['Ejecución material (partidas con precio)']}+{A_}{rows['Imprevistos']}")
ws.cell(row=rows["Total obra con IVA"], column=amount_col,
        value=f"={A_}{rows['Total obra sin IVA']}*(1+{A_}{rows['IVA (%)']})")
ws.cell(row=rows["Imprevistos (%)"], column=3).comment = Comment("Propuesta inicial. Ajustar según estado de la vivienda.", "Enzyme")
ws.cell(row=rows["IVA (%)"], column=3).comment = Comment("10 % para reformas de vivienda que cumplen los requisitos del tipo reducido; 21 % en otro caso. Confirmar en cada obra.", "Enzyme")
ws.cell(row=rows["Partidas con cantidad y sin precio"], column=3).comment = Comment("Partidas que entran en la oferta pero aún no tienen precio Enzyme: el total está incompleto mientras no sea 0.", "Enzyme")
ws.conditional_formatting.add(
    f"{P_}{first_line}:{P_}{last_line}",
    FormulaRule(formula=[f'AND({Q_}{first_line}>0,{P_}{first_line}="")'], fill=YELLOW))
ws.freeze_panes = ws.cell(row=HROW + 1, column=FIRST_ROOM_COL)
style_all(ws)

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(OUT)
